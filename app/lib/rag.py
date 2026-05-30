import os
import logging
from functools import lru_cache
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse
from typing import List, Optional
from langchain_core.documents import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from app.lib.config import settings

logger = logging.getLogger(__name__)


def _sync_database_url(url: str) -> str:
    """Convert the async application database URL to a sync PGVector connection string."""
    if not url:
        return url

    parsed = urlparse(url)
    scheme = parsed.scheme
    if scheme.endswith("+asyncpg"):
        scheme = scheme.replace("+asyncpg", "+psycopg2")
    elif scheme == "postgresql":
        scheme = "postgresql+psycopg2"

    query = parse_qs(parsed.query)
    incompatible_params = [
        "sslmode",
        "sslrootcert",
        "sslcert",
        "sslkey",
        "target_session_attrs",
        "channel_binding",
        "application_name",
        "connect_timeout",
    ]

    for param in incompatible_params:
        query.pop(param, None)

    return urlunparse(parsed._replace(scheme=scheme, query=urlencode(query, doseq=True)))


@lru_cache(maxsize=1)
def _get_embeddings():
    """Lazy-load the community FastEmbed embeddings so app startup stays lightweight."""
    try:
        from langchain_community.embeddings import FastEmbedEmbeddings
    except ImportError as exc:
        raise ImportError(
            "FastEmbedEmbeddings is unavailable. Install a compatible langchain-community "
            "package together with fastembed."
        ) from exc

    return FastEmbedEmbeddings(model_name=settings.EMBEDDING_MODEL)


def _get_vector_store_class():
    """Lazy-load the community PGVector store."""
    try:
        from langchain_community.vectorstores import PGVector
    except ImportError as exc:
        raise ImportError(
            "PGVector is unavailable. Install a compatible langchain-community package."
        ) from exc

    return PGVector

# Text splitter for chunking documents
text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=200,
    length_function=len,
    separators=["\n\n", "\n", ". ", " ", ""],
)


def load_file_documents(file_path: str, filename: str) -> List[Document]:
    """Load documents from a file (PDF, TXT, DOCX, CSV)."""
    ext = os.path.splitext(filename)[1].lower()
    documents = []
    
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    documents.append(Document(
                        page_content=text.strip(),
                        metadata={"source": filename, "page": i + 1, "type": "pdf"}
                    ))
        
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            if text.strip():
                documents.append(Document(
                    page_content=text.strip(),
                    metadata={"source": filename, "type": "txt"}
                ))
        
        elif ext == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            if text.strip():
                documents.append(Document(
                    page_content=text.strip(),
                    metadata={"source": filename, "type": "docx"}
                ))
        
        elif ext == ".csv":
            import csv
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader_csv = csv.reader(f)
                rows = list(reader_csv)
            if rows:
                header = rows[0] if rows else []
                for i, row in enumerate(rows[1:], start=2):
                    text = " | ".join(f"{h}: {v}" for h, v in zip(header, row) if v.strip())
                    if text:
                        documents.append(Document(
                            page_content=text,
                            metadata={"source": filename, "row": i, "type": "csv"}
                        ))
        
        elif ext == ".md":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            if text.strip():
                documents.append(Document(
                    page_content=text.strip(),
                    metadata={"source": filename, "type": "markdown"}
                ))
        
        else:
            # Try to read as plain text
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            if text.strip():
                documents.append(Document(
                    page_content=text.strip(),
                    metadata={"source": filename, "type": "unknown"}
                ))
    
    except Exception as e:
        logger.error(f"Error loading file {filename}: {e}")
        raise
    
    return documents


class RAGManager:
    """Manages Vector Storage and Retrieval for Store-specific data."""
    
    def __init__(self, store_id: str):
        self.store_id = store_id
        self.collection_name = f"store_{store_id.replace('-', '_')}"
        self.connection_string = _sync_database_url(settings.DATABASE_URL)
        self.vector_store = _get_vector_store_class()(
            connection_string=self.connection_string,
            embedding_function=_get_embeddings(),
            collection_name=self.collection_name,
            pre_delete_collection=False,
        )

    async def is_indexed(self) -> bool:
        """Check if the vector store already has data."""
        try:
            with self.vector_store._make_session() as session:
                collection = self.vector_store.get_collection(session)
                return bool(collection and collection.embeddings)
        except Exception:
            return False

    async def index_store_data(self, description: str, products: List[dict], force: bool = False):
        """Index store description and products into the vector store."""
        if not force and await self.is_indexed():
            # Already indexed, skip to avoid redundant processing
            return

        documents = []
        
        # Add store description
        if description:
            documents.append(Document(
                page_content=f"Store Description: {description}",
                metadata={"source": "description", "store_id": self.store_id}
            ))
            
        # Add products
        for product in products:
            content = f"Product: {product.get('name')} | Details: {product.get('description')} | Price: ${product.get('price')}"
            documents.append(Document(
                page_content=content,
                metadata={
                    "source": "product", 
                    "product_id": str(product.get('id', '')),
                    "product_name": product.get('name', ''),
                    "store_id": self.store_id
                }
            ))
            
        if documents:
            # Chunk if needed and add
            chunks = text_splitter.split_documents(documents)
            self.vector_store.add_documents(chunks)
            logger.info(f"Indexed {len(chunks)} chunks for store {self.store_id}")

    async def index_file(self, file_path: str, filename: str) -> int:
        """Index a file (PDF, TXT, DOCX, CSV, MD) into the vector store.
        
        Returns the number of chunks indexed.
        """
        # Load documents from the file
        raw_docs = load_file_documents(file_path, filename)
        
        if not raw_docs:
            logger.warning(f"No content extracted from {filename}")
            return 0
        
        # Add store_id metadata
        for doc in raw_docs:
            doc.metadata["store_id"] = self.store_id
            doc.metadata["filename"] = filename
        
        # Chunk the documents
        chunks = text_splitter.split_documents(raw_docs)
        
        if chunks:
            self.vector_store.add_documents(chunks)
            logger.info(f"Indexed {len(chunks)} chunks from file {filename} for store {self.store_id}")
        
        return len(chunks)

    async def query(self, text: str, k: int = 5) -> str:
        """Perform semantic search for relevant context."""
        results = self.vector_store.similarity_search(text, k=k)
        if not results:
            return ""
            
        context = "\n".join([res.page_content for res in results])
        return f"Relevant Information found in Knowledge Base:\n{context}"

    async def clear(self):
        """Clear all data from this store's vector store."""
        try:
            self.vector_store.delete_collection()
            logger.info(f"Cleared vector store for store {self.store_id}")
        except Exception as e:
            logger.error(f"Error clearing vector store: {e}")


# Helper to get manager instance
def get_rag_manager(store_id: str) -> RAGManager:
    return RAGManager(store_id)
